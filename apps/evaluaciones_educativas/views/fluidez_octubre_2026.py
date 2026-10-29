from apps.evaluaciones_educativas.models.fluidez_octubre_2026 import *
from apps.consultasge.models import CapaUnicaOfertas
from apps.evaluaciones_educativas.forms.fluidez_octubre_2026 import *
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse, Http404
from django.core.paginator import Paginator
from django.db import transaction
from django.contrib.auth.decorators import login_required
from django.templatetags.static import static
from django.db.models import Count, Q, Avg, F
from datetime import date, datetime
import json
from openpyxl import Workbook
from django.core.exceptions import PermissionDenied
from django.contrib import messages
from django.urls import reverse
from urllib.parse import urlencode
import re
import logging
from django.db import DatabaseError, IntegrityError, connections
import uuid
from django.views.decorators.http import require_POST
from apps.evaluaciones_educativas.services.fluidez_octubre_2026 import (
    secciones_region,
    alumnos_seccion_grado,
    GRADOS_URL,
    MAX_PALABRAS,
)

logger = logging.getLogger(__name__)


# def actualizar_seccion(request, alumno_public_id):
#     """Actualiza la sección, turno, comunidad_indigena y discapacidad de un alumno."""
#     if request.method == "POST":
#         fid = request.POST.get("fid")

#         alumno = get_object_or_404(AlumnoFluidez2026, public_id=alumno_public_id)

#         nueva_seccion_turno_public_id = request.POST.get("seccion_turno")
#         comunidad = request.POST.get("comunidad_indigena")
#         discapacidad = request.POST.get("discapacidad")

#         if nueva_seccion_turno_public_id:
#             seccion_obj = get_object_or_404(
#                 SeccionFluidez2026, public_id=nueva_seccion_turno_public_id
#             )
#             alumno.seccion = seccion_obj
#         else:
#             alumno.seccion = None

#         valores_comunidad = [
#             str(v) for v, _ in AlumnoFluidez2026.OPCIONES_COMUNIDAD_INDIGENA
#         ]
#         if comunidad in valores_comunidad:
#             alumno.comunidad_indigena = comunidad
#         elif not comunidad:
#             alumno.comunidad_indigena = None

#         valores_discapacidad = [
#             str(v) for v, _ in AlumnoFluidez2026.OPCIONES_DISCAPACIDAD
#         ]
#         if discapacidad in valores_discapacidad:
#             alumno.discapacidad = discapacidad
#         elif not discapacidad:
#             alumno.discapacidad = None

#         alumno.save()

#         base_url = reverse("evaluaciones_educativas:fluidez_octubre_2026:lista")

#         if fid:
#             return redirect(f"{base_url}?fid={fid}")

#         return redirect(base_url)

#     return redirect("evaluaciones_educativas:fluidez_octubre_2026:lista")


# def carga_alumno(request, fid_actual, grado_public_id):
#     cuil = 20422632264
#     grado = get_object_or_404(GradoFluidez2026, public_id=grado_public_id)
#     alumno_form = AlumnoFluidezOctubre2026Form()
#     grado_form = GradoFluidezOctubre2026Form(instance=grado)
#     seccion_form = SeccionFluidezOctubre2026Form(
#         cueanexo=grado.cueanexo, nombre_grado=grado.nombre_grado
#     )
#     if request.method == "POST":
#         alumno_form = AlumnoFluidezOctubre2026Form(request.POST)
#         seccion_form = SeccionFluidezOctubre2026Form(
#             request.POST, cueanexo=grado.cueanexo, nombre_grado=grado.nombre_grado
#         )
#         if alumno_form.is_valid() and seccion_form.is_valid():
#             with transaction.atomic():
#                 seccion_turno_id = seccion_form.cleaned_data["seccion_turno"]
#                 instancia_seccion, creado_seccion = (
#                     SeccionFluidez2026.objects.get_or_create(
#                         id=seccion_turno_id, grado=grado
#                     )
#                 )
#                 alumno = alumno_form.save(commit=False)
#                 alumno.seccion = instancia_seccion
#                 alumno.save()
#             return redirect(
#                 "evaluaciones_educativas:fluidez_octubre_2026:lista", fid_actual=fid_actual
#             )

#     context = {
#         "alumno_form": alumno_form,
#         "grado_form": grado_form,
#         "seccion_form": seccion_form,
#         "fid_actual": fid_actual,
#     }
#     return render(request, "fluidez_octubre_2026/alumno.html", context)


# def editar_alumno(request, alumno_public_id, fid_actual):
#     instancia_alumno = get_object_or_404(AlumnoFluidez2026, public_id=alumno_public_id)
#     if not instancia_alumno.seccion.grado.cueanexo:
#         alumno_datos = get_object_or_404(
#             TablaTemporalAlumnoFluidez2026, numero_de_documento=instancia_alumno.dni
#         )
#         cueanexo = alumno_datos.cueanexo
#         anio = alumno_datos.anio
#     cueanexo = instancia_alumno.seccion.grado.cueanexo
#     anio = instancia_alumno.seccion.grado.nombre_grado
#     seccion = instancia_alumno.seccion.id

#     instancia_grado = get_object_or_404(
#         GradoFluidez2026, cueanexo=cueanexo, nombre_grado=anio
#     )
#     alumno_form = AlumnoFluidezOctubre2026Form(instance=instancia_alumno)
#     seccion_form = SeccionFluidezOctubre2026Form(
#         cueanexo=instancia_grado.cueanexo,
#         nombre_grado=instancia_grado.nombre_grado,
#         initial={"seccion_turno": seccion},
#     )
#     grado_form = GradoFluidezOctubre2026Form(instance=instancia_grado)
#     if request.method == "POST":
#         alumno_form = AlumnoFluidezOctubre2026Form(request.POST, instance=instancia_alumno)
#         grado_form = GradoFluidezOctubre2026Form(request.POST, instance=instancia_grado)
#         seccion_form = SeccionFluidezOctubre2026Form(
#             request.POST, cueanexo=cueanexo, nombre_grado=instancia_grado.nombre_grado
#         )
#         if alumno_form.is_valid() and grado_form.is_valid() and seccion_form.is_valid():
#             with transaction.atomic():
#                 instancia_grado, creado_grado = GradoFluidez2026.objects.get_or_create(
#                     nombre_grado=anio, cueanexo=cueanexo
#                 )
#                 seccion_turno_id = seccion_form.cleaned_data["seccion_turno"]
#                 instancia_seccion, creado_seccion = (
#                     SeccionFluidez2026.objects.get_or_create(
#                         id=seccion_turno_id, grado=instancia_grado
#                     )
#                 )
#                 alumno = alumno_form.save(commit=False)
#                 alumno.seccion = instancia_seccion
#                 alumno.save()
#             return redirect(
#                 "evaluaciones_educativas:fluidez_octubre_2026:lista", fid_actual=fid_actual
#             )
#     context = {
#         "alumno_form": alumno_form,
#         "grado_form": grado_form,
#         "seccion_form": seccion_form,
#         "grado_public": instancia_grado.public_id,
#         "fid_actual": fid_actual,
#     }
#     return render(request, "fluidez_octubre_2026/alumno.html", context)


# def borrar_registro_alumno(request, alumno_public_id, fid_actual):
#     alumno_id = get_object_or_404(AlumnoFluidez2026, public_id=alumno_public_id)
#     if request.method == "POST":
#         form = BorrarRegistroAlumnoOctubre2026Form(request.POST)
#         if form.is_valid():
#             with transaction.atomic():
#                 eleccion = form.cleaned_data["borrar"]
#                 if eleccion:
#                     if alumno_id.dni:
#                         if TablaTemporalAlumnoFluidez2026.objects.filter(
#                             numero_de_documento=alumno_id.dni
#                         ).exists():
#                             TablaTemporalAlumnoFluidez2026.objects.filter(
#                                 numero_de_documento=alumno_id.dni
#                             ).delete()
#                     alumno_id.delete()
#             return redirect("evaluaciones_educativas:fluidez_octubre_2026:monitoreo_alumno")
#     else:
#         form = BorrarRegistroAlumnoOctubre2026Form()
#     context = {"form": form, "alumno": alumno_id, "fid_actual": fid_actual}
#     return render(request, "fluidez_octubre_2026/borrar_registro_alumno.html", context)


# def monitoreo(request):
#     filtro_escuela = request.GET.get("escuela", "").strip()
#     filtro_sector = request.GET.get("sector", "").strip()
#     filtro_ambito = request.GET.get("ambito", "").strip()
#     filtro_region = request.GET.get("region", "").strip()
#     filtro_cueanexo = request.GET.get("cueanexo", "").strip()

#     queryset = EstablecimientosFluidez2026.objects.prefetch_related(
#         "gradofluidez2026_set"
#     )

#     if filtro_escuela:
#         queryset = queryset.filter(escuela__icontains=filtro_escuela)
#     if filtro_sector:
#         queryset = queryset.filter(sector=filtro_sector)
#     if filtro_ambito:
#         queryset = queryset.filter(ambito=filtro_ambito)
#     if filtro_region:
#         queryset = queryset.filter(region=filtro_region)
#     if filtro_cueanexo:
#         queryset = queryset.filter(cueanexo__icontains=filtro_cueanexo)

#     paginator = Paginator(queryset.order_by("escuela"), 100)
#     page_number = request.GET.get("page")
#     page_obj = paginator.get_page(page_number)

#     query_params = request.GET.copy()
#     if "page" in query_params:
#         del query_params["page"]
#     params_url = query_params.urlencode()

#     cueanexos_pagina = [est.cueanexo for est in page_obj]

#     temporal_qs = TablaTemporalAlumnoFluidez2026.objects.filter(
#         cueanexo__in=cueanexos_pagina
#     ).values("cueanexo", "anio", "numero_de_documento")

#     temporal_dict = {}
#     for item in temporal_qs:
#         k = (str(item["cueanexo"]), item["anio"])
#         if k not in temporal_dict:
#             temporal_dict[k] = set()
#         if item["numero_de_documento"]:
#             temporal_dict[k].add(item["numero_de_documento"])

#     alumnos_qs = AlumnoFluidez2026.objects.filter(
#         seccion__grado__Establecimiento__cueanexo__in=cueanexos_pagina
#     ).values(
#         "dni",
#         "seccion__grado__cueanexo",
#         "seccion__grado__nombre_grado",
#         "fluidez_lectora_octubre_2026__evaluacion_id",
#     )

#     alumnos_por_grado = {}
#     for al in alumnos_qs:
#         k = (al["seccion__grado__cueanexo"], al["seccion__grado__nombre_grado"])
#         if k not in alumnos_por_grado:
#             alumnos_por_grado[k] = []
#         alumnos_por_grado[k].append(al)

#     for est in page_obj:
#         for grado in est.gradofluidez2026_set.all():
#             cueanexo = est.cueanexo
#             nombre_grado = grado.nombre_grado
#             anio_temporal = (
#                 "2do Grado/Año" if nombre_grado == "2do Año/Grado" else "3er Grado/Año"
#             )
#             dnis_temporales = temporal_dict.get((cueanexo, anio_temporal), set())
#             grado.alumnos_esperados = len(dnis_temporales)
#             alumnos_del_grado = alumnos_por_grado.get((cueanexo, nombre_grado), [])

#             dnis_extras = {
#                 al["dni"]
#                 for al in alumnos_del_grado
#                 if al["dni"] and al["dni"] not in dnis_temporales
#             }

#             dnis_interes = dnis_temporales.union(dnis_extras)
#             alumnos_existentes = [
#                 al for al in alumnos_del_grado if al["dni"] in dnis_interes
#             ]

#             grado.alumnos_esperados = len(dnis_temporales)
#             grado.total_alumnos = len(alumnos_existentes)
#             grado.total_evaluados = sum(
#                 1
#                 for al in alumnos_existentes
#                 if al["fluidez_lectora_octubre_2026__evaluacion_id"] is not None
#             )

#     sectores = (
#         EstablecimientosFluidez2026.objects.values_list("sector", flat=True)
#         .distinct()
#         .order_by("sector")
#     )
#     ambitos = (
#         EstablecimientosFluidez2026.objects.values_list("ambito", flat=True)
#         .distinct()
#         .order_by("ambito")
#     )
#     regiones = (
#         EstablecimientosFluidez2026.objects.values_list("region", flat=True)
#         .distinct()
#         .order_by("region")
#     )

#     contexto = {
#         "establecimientos": page_obj,
#         "page_obj": page_obj,
#         "params_url": params_url,
#         "sectores": sectores,
#         "ambitos": ambitos,
#         "regiones": regiones,
#         "valores_filtros": {
#             "escuela": filtro_escuela,
#             "sector": filtro_sector,
#             "ambito": filtro_ambito,
#             "region": filtro_region,
#             "cueanexo": filtro_cueanexo,
#         },
#     }
#     return render(request, "fluidez_octubre_2026/monitoreo.html", contexto)


# def monitoreo_alumno(request):
#     buscar = request.GET.get("buscar", "").strip()
#     alumnos_qs = None

#     if buscar:
#         buscar_limpio = buscar.replace(".", "").replace(" ", "")

#         lista_dnis = list(
#             TablaTemporalAlumnoFluidez2026.objects.filter(
#                 numero_de_documento=buscar_limpio
#             ).values_list("numero_de_documento", flat=True)
#         )

#         lista = list(
#             AlumnoFluidez2026.objects.filter(dni=buscar_limpio).values_list(
#                 "dni", flat=True
#             )
#         )

#         lista_dnis.extend(lista)
#         lista_dnis = list(set(filter(None, lista_dnis)))

#         alumnos_qs = list(
#             AlumnoFluidez2026.objects.filter(dni__in=lista_dnis).select_related(
#                 "seccion__grado__Establecimiento", "fluidez_lectora_octubre_2026"
#             )
#         )

#         temporal_dict = {
#             t.numero_de_documento: t.nombre_institucion
#             for t in TablaTemporalAlumnoFluidez2026.objects.filter(
#                 numero_de_documento__in=lista_dnis
#             ).only("numero_de_documento", "nombre_institucion")
#         }

#         for alumno in alumnos_qs:
#             alumno.establecimiento_original = temporal_dict.get(alumno.dni, "-")

#     contexto = {"alumnos_resultados": alumnos_qs}
#     return render(request, "fluidez_octubre_2026/monitoreo_alumno.html", contexto)


#-------------------------------------------------CARGA DE EXÁMENES (TABULADORES)-------------------------------------------------
# El tabulador es el usuario logueado (su CUIL está en TabuladoresFluidezOctubre2026)
# y ve los alumnos de 2º y 3º de las escuelas de su lista_cueanexos. Los alumnos,
# secciones y grados son de SGE y no se modifican: solo se cargan los exámenes.

# Namespace del lock pg_advisory_xact_lock(namespace, id_persona) al guardar un examen.
LOCK_EXAMEN_FLUIDEZ_OCT = 202610


def _tabulador_actual(request):
    return TabuladoresFluidezOctubre2026.objects.filter(cuil=str(request.user.username)).first()


def _cueanexos_tabulador(tabulador):
    return [str(c).strip() for c in (tabulador.lista_cueanexos or [])]


def _validar_cueanexo(tabulador, cueanexo):
    if str(cueanexo).strip() not in _cueanexos_tabulador(tabulador):
        raise PermissionDenied('La escuela no está asignada a este tabulador.')


def _anio_grado(grado):
    anio_grado = GRADOS_URL.get(str(grado))
    if not anio_grado:
        raise Http404('Grado inválido.')
    return anio_grado


def _sin_tabulador(request):
    messages.error(request, 'Tu usuario no está cargado como tabulador de Fluidez Octubre 2026.')
    return redirect(reverse('evaluaciones_educativas:dashboard'))


def _examenes_por_persona(ids_persona):
    """{id_persona: FluidezLectoraOctubre2026} de los alumnos que ya tienen examen."""
    if not ids_persona:
        return {}
    return {
        e.alumno_id: e
        for e in FluidezLectoraOctubre2026.objects.filter(
            evaluacion__operativo=FluidezLectoraOctubre2026.operativo(),
            alumno_id__in=ids_persona,
        )
    }


def _carga_cerrada(cueanexo, anio_grado):
    return CierreCargaFluidezOctubre2026.objects.filter(
        cueanexo=cueanexo, anio_grado=anio_grado, estado_carga=True
    ).exists()


def _listado_alumnos(request, fid_actual, template, con_examenes):
    """
    Lógica común de `lista` y `lista_examen`. El filtro (cueanexo y grado) se
    guarda en sesión por pestaña (`fid`, va en la URL), con POST -> redirect -> GET.
    """
    tabulador = _tabulador_actual(request)
    if not tabulador:
        return _sin_tabulador(request)

    fid_actual = fid_actual or request.GET.get('fid') or uuid.uuid4().hex[:12]
    clave = f'filtro_oct_{fid_actual}'

    if request.method == 'POST':
        cueanexo_nuevo = request.POST.get('cueanexo') or None
        grado_nuevo = request.POST.get('grado') or None
        # Si cambió la escuela, el grado elegido antes ya no corresponde.
        if cueanexo_nuevo != request.session.get(clave, {}).get('cueanexo'):
            grado_nuevo = None
        request.session[clave] = {'cueanexo': cueanexo_nuevo, 'grado': grado_nuevo}
        request.session.modified = True
        return redirect(request.path)

    cueanexos = _cueanexos_tabulador(tabulador)
    datos = request.session.get(clave) or {}
    cueanexo = datos.get('cueanexo')
    grado = datos.get('grado')
    if cueanexo not in cueanexos:
        cueanexo, grado = None, None
    if grado not in GRADOS_URL:
        grado = None

    cueanexo_form = CueanexoFluidezOctubre2026ViewForm({'cueanexo': cueanexo}, cueanexos=cueanexos)
    grado_form = GradoFluidezOctubre2026ViewForm({'grado': grado}) if cueanexo else None

    alumnos = None
    anio_grado = GRADOS_URL.get(grado) if grado else None
    carga_cerrada = False
    if cueanexo and anio_grado:
        try:
            alumnos = alumnos_seccion_grado(cueanexo, anio_grado)
        except DatabaseError:
            logger.exception('Error consultando alumnos en sge_nacion (%s, %s)', cueanexo, anio_grado)
            messages.error(request, ERROR_SGE)
            alumnos = []
        if con_examenes:
            examenes = _examenes_por_persona([a['id_persona'] for a in alumnos])
            for a in alumnos:
                a['examen'] = examenes.get(a['id_persona'])
            carga_cerrada = _carga_cerrada(cueanexo, anio_grado)

    contexto = {
        'tabulador': tabulador,
        'sin_escuelas': not cueanexos,
        'cueanexo_form': cueanexo_form,
        'grado_form': grado_form,
        'cueanexo': cueanexo,
        'grado': grado,
        'anio_grado': anio_grado,
        'alumnos': alumnos,
        'carga_cerrada': carga_cerrada,
        'fid_actual': fid_actual,
    }
    return render(request, template, contexto)


@login_required
def lista(request, fid_actual=None):
    """Listado de solo lectura de los alumnos de SGE."""
    return _listado_alumnos(request, fid_actual, 'fluidez_octubre_2026/lista.html', con_examenes=False)


@login_required
def lista_examen(request, fid_actual=None):
    """Alumnos con el estado de su examen, para cargarlo o editarlo."""
    return _listado_alumnos(request, fid_actual, 'fluidez_octubre_2026/lista_examen.html', con_examenes=True)


@login_required
def evaluacion(request, fid_actual, cueanexo, grado, id_persona):
    """Alta o edición del examen de un alumno."""
    tabulador = _tabulador_actual(request)
    if not tabulador:
        return _sin_tabulador(request)
    _validar_cueanexo(tabulador, cueanexo)
    anio_grado = _anio_grado(grado)
    url_lista = reverse(
        'evaluaciones_educativas:fluidez_octubre_2026:lista_examen_fid',
        kwargs={'fid_actual': fid_actual},
    )

    try:
        encontrados = alumnos_seccion_grado(cueanexo, anio_grado, id_persona=id_persona)
    except DatabaseError:
        logger.exception('Error consultando el alumno %s en sge_nacion', id_persona)
        messages.error(request, ERROR_SGE)
        return redirect(url_lista)
    if not encontrados:
        raise Http404('El alumno no pertenece a esa escuela y grado.')
    alumno = encontrados[0]

    if _carga_cerrada(cueanexo, anio_grado):
        messages.error(request, 'La carga de este grado está cerrada. Abrila para modificar exámenes.')
        return redirect(url_lista)

    existente = FluidezLectoraOctubre2026.de_alumno(id_persona)
    form = EvaluacionFluidezOctubre2026Form(
        request.POST or None, instance=existente, max_cantidad_palabra=MAX_PALABRAS[anio_grado],
    )

    if request.method == 'POST' and form.is_valid():
        db = FluidezLectoraOctubre2026.objects.db
        with transaction.atomic(using=db):
            # Serializa los guardados del mismo alumno (doble envío, dos pestañas):
            # se vuelve a buscar el examen ya con el lock tomado.
            with connections[db].cursor() as cur:
                cur.execute('SELECT pg_advisory_xact_lock(%s, %s)', [LOCK_EXAMEN_FLUIDEZ_OCT, id_persona])
            examen = FluidezLectoraOctubre2026.de_alumno(id_persona)
            if examen:
                for campo, valor in form.cleaned_data.items():
                    setattr(examen, campo, valor)
                examen.encargado_carga = 'TABULADOR'
                examen.alumno_id = id_persona
                examen.save()
            else:
                FluidezLectoraOctubre2026.crear(
                    id_alumno=id_persona,
                    cueanexo_original=alumno['cueanexo'],
                    cueanexo=str(cueanexo).strip(),
                    alumno_id=id_persona,
                    encargado_carga='TABULADOR',
                    **form.cleaned_data,
                )
        messages.success(request, f"Examen de {alumno['apellido']}, {alumno['nombre']} guardado correctamente.")
        return redirect(url_lista)

    context = {
        'form': form,
        'alumno': alumno,
        'anio_grado': anio_grado,
        'es_edicion': bool(existente),
        'fid_actual': fid_actual,
    }
    return render(request, 'fluidez_octubre_2026/evaluacion.html', context)


@login_required
@require_POST
def completar_carga(request, cueanexo, grado):
    """
    Abre o cierra la carga de un grado. Cerrar solo se permite si todos los
    alumnos tienen examen; abrir, siempre.
    """
    tabulador = _tabulador_actual(request)
    if not tabulador:
        return JsonResponse({'error': 'Sin permiso.'}, status=403)
    _validar_cueanexo(tabulador, cueanexo)
    anio_grado = _anio_grado(grado)

    cierre, _ = CierreCargaFluidezOctubre2026.objects.get_or_create(cueanexo=cueanexo, anio_grado=anio_grado)
    faltantes = None
    es_valido = True

    if cierre.estado_carga:
        cierre.estado_carga = False
    else:
        try:
            ids = [a['id_persona'] for a in alumnos_seccion_grado(cueanexo, anio_grado)]
        except DatabaseError:
            logger.exception('Error consultando alumnos en sge_nacion (%s, %s)', cueanexo, anio_grado)
            return JsonResponse({'error': ERROR_SGE}, status=503)
        faltantes = len(ids) - len(_examenes_por_persona(ids))
        if faltantes == 0:
            cierre.estado_carga = True
        else:
            es_valido = False

    if es_valido:
        cierre.cuil_tabulador = tabulador.cuil
        cierre.save()

    return JsonResponse({
        'es_valido': es_valido,
        'nuevo_estado': cierre.estado_carga,
        'faltantes': faltantes,
    })


@login_required
def descargar_excel(request, cueanexo, grado):
    """Excel con todos los alumnos del grado y su examen ('SIN CARGAR' si no tiene)."""
    tabulador = _tabulador_actual(request)
    if not tabulador:
        return _sin_tabulador(request)
    _validar_cueanexo(tabulador, cueanexo)
    anio_grado = _anio_grado(grado)

    try:
        alumnos = alumnos_seccion_grado(cueanexo, anio_grado)
    except DatabaseError:
        logger.exception('Error consultando alumnos en sge_nacion (%s, %s)', cueanexo, anio_grado)
        messages.error(request, ERROR_SGE)
        return redirect(reverse('evaluaciones_educativas:fluidez_octubre_2026:lista_examen'))
    examenes = _examenes_por_persona([a['id_persona'] for a in alumnos])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Fluidez Octubre 2026'
    ws['A1'] = f'CUEANEXO: {cueanexo}'
    ws['G1'] = f'FECHA Y HORA: {datetime.now():%d/%m/%Y %H:%M:%S}'
    ws.append([
        'NOMBRE', 'APELLIDO', 'DNI', 'COMUNIDAD INDÍGENA', 'DISCAPACIDAD', 'GRADO',
        'SECCIÓN', 'TURNO', 'ASISTENCIA', 'FLUIDEZ', 'P1', 'P2', 'P3', 'P4', 'P5', 'P6',
    ])
    for a in alumnos:
        e = examenes.get(a['id_persona'])
        fila = [
            a['nombre'], a['apellido'], a['numero_documento'], a['comunidad_indigena'],
            a['discapacidad'], a['anio_grado'], str(a['nombre_seccion'] or ''), a['turno'],
        ]
        if e:
            fila += [
                e.asistencia, e.cantidad_palabras_leidas, e.pregunta_1, e.pregunta_2,
                e.pregunta_3, e.pregunta_4, e.pregunta_5, e.pregunta_6,
            ]
        else:
            fila.append('SIN CARGAR')
        ws.append(fila)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = (
        f'attachment; filename="fluidez_octubre_2026_{cueanexo}_{grado}_grado_{datetime.now():%Y%m%d_%H%M}.xlsx"'
    )
    wb.save(response)
    return response


#-------------------------------------------------CARGA Y SELECCION DE TABULADOR-------------------------------------------------------------


@login_required
def gestion_tabuladores(request):
    """Página central de gestión de tabuladores para el rol Regional."""
    nivel_acceso = request.user.nivelacceso_id
    if nivel_acceso != "Regional":
        messages.error(request, 'No tienes permiso para acceder a esta página.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    cuil = str(request.user.username)
    regiones = obtener_regional(cuil)

    # Si el usuario no tiene ninguna región asignada
    if not regiones:
        messages.error(request, 'No se encontró una región asignada a tu usuario.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))
    #print(regiones)
    if 'R.E. 2' in regiones:
        regiones = ['R.E. 2','SUB. R.E. 2-B']

    # ── Selector de región ──────────────────────────────────────────
    # Si viene ?region=X en la URL y es una región válida del usuario → guardar en sesión
    region_solicitada = request.GET.get('region')
    if region_solicitada and region_solicitada in regiones:
        request.session['tab_fluidez_oct26_region'] = region_solicitada

    # Leer la región activa desde sesión; si no hay o ya no es válida, usar la primera
    region = request.session.get('tab_fluidez_oct26_region')
    if not region or region not in regiones:
        region = regiones[0]
        request.session['tab_fluidez_oct26_region'] = region

    # Cupo habilitado para esta región
    cupo_obj = TemporalCargaTabuladoresFluidez2026Eliminar.objects.filter(region=region).first()
    tabuladores_permitidos = cupo_obj.tabuladores_permitidos if cupo_obj else 0

    # Cuántos ya fueron cargados para esta región
    tabuladores_cargados = TabuladoresFluidezOctubre2026.objects.filter(region=region).count()

    puede_cargar = tabuladores_permitidos > 0 and tabuladores_cargados < tabuladores_permitidos

    # Lista de tabuladores ya cargados para mostrar en la tabla
    tabuladores_list = TabuladoresFluidezOctubre2026.objects.filter(region=region).order_by('apellido', 'nombre')

    context = {
        'region':                 region,
        'regiones':               regiones,          # todas las regiones del usuario
        'multiples_regiones':     len(regiones) > 1,
        'tabuladores_permitidos': tabuladores_permitidos,
        'tabuladores_cargados':   tabuladores_cargados,
        'puede_cargar':           puede_cargar,
        'tabuladores_list':       tabuladores_list,
    }
    return render(request, 'fluidez_octubre_2026/gestion_tabuladores.html', context)


@login_required
def carga_tabulador(request):
    """Formulario para cargar un nuevo tabulador (solo Regional con cupo disponible)."""
    nivel_acceso = request.user.nivelacceso_id
    if nivel_acceso != "Regional":
        messages.error(request, 'No tienes permiso para acceder a esta página.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    cuil = str(request.user.username)
    regiones = obtener_regional(cuil)
    #print(regiones)

    if not regiones:
        messages.error(request, 'No se encontró una región asignada a tu usuario.')
        return redirect(reverse('evaluaciones_educativas:fluidez_octubre_2026:gestion_tabuladores'))
    
    if 'R.E. 2' in regiones:
        regiones = ['R.E. 2','SUB. R.E. 2-B']

    # Leer la región activa desde sesión (coherente con gestion_tabuladores)
    region = request.session.get('tab_fluidez_oct26_region')
    if not region or region not in regiones:
        region = regiones[0]
        request.session['tab_fluidez_oct26_region'] = region

    # Verificar cupo
    cupo_obj = TemporalCargaTabuladoresFluidez2026Eliminar.objects.filter(region=region).first()
    tabuladores_permitidos = cupo_obj.tabuladores_permitidos if cupo_obj else 0
    tabuladores_cargados   = TabuladoresFluidezOctubre2026.objects.filter(region=region).count()

    if tabuladores_permitidos == 0 or tabuladores_cargados >= tabuladores_permitidos:
        messages.error(request, f'No podés cargar más tabuladores. La región {region} ya tiene el cupo completo ({tabuladores_cargados}/{tabuladores_permitidos}).')
        return redirect(reverse('evaluaciones_educativas:fluidez_octubre_2026:gestion_tabuladores'))

    if request.method == 'POST':
        form = TabuladorFluidezOctubreForm(request.POST)
        if form.is_valid():
            tabulador = form.save(commit=False)
            tabulador.region = region          # asignar la región automáticamente
            tabulador.lista_cueanexos = []     # inicializar vacío, se completa en el 2do botón
            tabulador.save()
            messages.success(request, f'Tabulador {tabulador.apellido}, {tabulador.nombre} (CUIL: {tabulador.cuil}) cargado correctamente.')
            return redirect(reverse('evaluaciones_educativas:fluidez_octubre_2026:gestion_tabuladores'))
    else:
        form = TabuladorFluidezOctubreForm()

    context = {
        'form':                   form,
        'region':                 region,
        'tabuladores_permitidos': tabuladores_permitidos,
        'tabuladores_cargados':   tabuladores_cargados,
        'cupo_restante':          tabuladores_permitidos - tabuladores_cargados,
    }
    return render(request, 'fluidez_octubre_2026/carga_tabulador.html', context)



@login_required
def eliminar_tabulador(request, cuil_tabulador):
    """Elimina un tabulador cargado (solo Regional, y solo de su propia región)."""
    nivel_acceso = request.user.nivelacceso_id
    if nivel_acceso != "Regional":
        messages.error(request, 'No tienes permiso para acceder a esta página.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))
    #--- ES REDUNDANDTE PORQUE YA ESTAMOS CON UN USUARIO EL CUAL TIENE UNA REGION   
    cuil_usuario = str(request.user.username)
    regiones = obtener_regional(cuil_usuario)
    if not regiones:
        messages.error(request, 'No se encontró una región asignada a tu usuario.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))
    if 'R.E. 2' in regiones:
        regiones = ['R.E. 2','SUB. R.E. 2-B']
    region = request.session.get('tab_fluidez_oct26_region')
    if not region or region not in regiones:
        region = regiones[0]
        request.session['tab_fluidez_oct26_region'] = region

     #--- ES REDUNDANDTE PORQUE YA ESTAMOS CON UN USUARIO EL CUAL TIENE UNA REGION----FIN---  

    if request.method == 'POST':
        tabulador = TabuladoresFluidezOctubre2026.objects.filter(cuil=cuil_tabulador, region=region).first()
        if tabulador:
            nombre_completo = f'{tabulador.apellido}, {tabulador.nombre}'
            tabulador.delete()
            messages.success(request, f'Tabulador {nombre_completo} (CUIL: {cuil_tabulador}) eliminado correctamente.')
        else:
            messages.warning(request, 'El tabulador ya no existe.')

    return redirect(reverse('evaluaciones_educativas:fluidez_octubre_2026:gestion_tabuladores'))


@login_required
def asignacion_tabulador(request, cuil_tabulador):
    """Vista para asignar/quitar escuelas a un tabulador específico."""
    nivel_acceso = request.user.nivelacceso_id
    if nivel_acceso != "Regional":
        messages.error(request, 'No tienes permiso para acceder a esta página.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    # Obtener región activa desde sesión
    cuil_usuario = str(request.user.username)
    regiones = obtener_regional(cuil_usuario)
    if not regiones:
        messages.error(request, 'No se encontró una región asignada a tu usuario.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))
    
    if 'R.E. 2' in regiones:
        regiones = ['R.E. 2','SUB. R.E. 2-B']
    region = request.session.get('tab_fluidez_oct26_region')
    if not region or region not in regiones:
        region = regiones[0]
        request.session['tab_fluidez_oct26_region'] = region

    # Obtener el tabulador — debe pertenecer a la región activa
    tabulador = get_object_or_404(TabuladoresFluidezOctubre2026, cuil=cuil_tabulador, region=region)

    # ── POST: agregar o quitar una escuela ────────────────────────────
    if request.method == 'POST':
        accion   = request.POST.get('accion')
        cueanexo = request.POST.get('cueanexo', '').strip()

        nivel, texto = 'error', 'Acción no válida.'

        if accion == 'agregar' and cueanexo:
            ya_asignada = TabuladoresFluidezOctubre2026.objects.filter(
                region=region,
                lista_cueanexos__contains=[cueanexo]
            ).exclude(cuil=cuil_tabulador).exists()

            if ya_asignada:
                nivel, texto = 'error', f'La escuela {cueanexo} ya está asignada a otro tabulador de esta región.'
            else:
                lista = list(tabulador.lista_cueanexos or [])
                if cueanexo not in lista:
                    lista.append(cueanexo)
                    tabulador.lista_cueanexos = lista
                    tabulador.save(update_fields=['lista_cueanexos'])
                    nivel, texto = 'success', f'Escuela {cueanexo} asignada correctamente.'
                else:
                    nivel, texto = 'warning', f'La escuela {cueanexo} ya estaba en la lista.'

        elif accion == 'quitar' and cueanexo:
            lista = list(tabulador.lista_cueanexos or [])
            if cueanexo in lista:
                lista.remove(cueanexo)
                tabulador.lista_cueanexos = lista
                tabulador.save(update_fields=['lista_cueanexos'])
                nivel, texto = 'success', f'Escuela {cueanexo} quitada correctamente.'
            else:
                nivel, texto = 'warning', f'La escuela {cueanexo} no estaba en la lista.'

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'ok': nivel in ('success', 'warning'),
                'nivel': nivel,
                'mensaje': texto,
                'cueanexo': cueanexo,
                'accion': accion,
            }, status=200 if nivel != 'error' else 409)

        getattr(messages, nivel)(request, texto)
        return redirect(reverse(
            'evaluaciones_educativas:fluidez_octubre_2026:asignacion_tabulador',
            kwargs={'cuil_tabulador': cuil_tabulador}
        ))

    # ── GET: construir listas de escuelas ────────────────────────────
    # Refrescar el tabulador desde DB para tener el estado más reciente
    tabulador.refresh_from_db()
    cueanexos_propios = [str(c) for c in (tabulador.lista_cueanexos or [])]

    #print(f"[asignacion] GET region={region} cueanexos_propios={cueanexos_propios}")

    # Todos los cueanexos ocupados en la región (por cualquier tabulador)
    cueanexos_ocupados = set()
    for tab in TabuladoresFluidezOctubre2026.objects.filter(region=region):
        for c in (tab.lista_cueanexos or []):
            cueanexos_ocupados.add(str(c))

    # Escuelas disponibles en la región con la oferta requerida
    todas_escuelas = list(
        CapaUnicaOfertas.objects.filter(
            region_loc=region,
            oferta__icontains='Común - Primaria de 7 años'
        ).values('cueanexo', 'nom_est', 'localidad').distinct().order_by('nom_est')
    )
    # Normalizar cueanexo a string: la columna puede devolver int/Decimal
    # según el driver de DB, mientras que lista_cueanexos siempre guarda strings.
    for e in todas_escuelas:
        e['cueanexo'] = str(e['cueanexo']).strip()

    #print(f"[asignacion] todas_escuelas count={len(todas_escuelas)}")

    # ── Escuelas asignadas: buscar directamente por cueanexo (sin filtro región)
    # para no depender de que region_loc coincida exactamente.
    if cueanexos_propios:
        escuelas_asignadas_qs = list(
            CapaUnicaOfertas.objects.filter(
                cueanexo__in=cueanexos_propios,
                oferta__icontains='Común - Primaria de 7 años'
            ).values('cueanexo', 'nom_est', 'localidad').distinct().order_by('nom_est')
        )
        for e in escuelas_asignadas_qs:
            e['cueanexo'] = str(e['cueanexo']).strip()
        # Incluir cueanexos que no se encuentren en la tabla (como fallback)
        cueanexos_encontrados = {e['cueanexo'] for e in escuelas_asignadas_qs}
        for c in cueanexos_propios:
            if c not in cueanexos_encontrados:
                escuelas_asignadas_qs.append({'cueanexo': c, 'nom_est': f'(Escuela {c})', 'localidad': ''})
        escuelas_asignadas = escuelas_asignadas_qs
    else:
        escuelas_asignadas = []

    # Disponibles: las que no están ocupadas por ningún tabulador
    escuelas_disponibles = [e for e in todas_escuelas if e['cueanexo'] not in cueanexos_ocupados]

    #print(f"[asignacion] asignadas={len(escuelas_asignadas)} disponibles={len(escuelas_disponibles)}")

    context = {
        'tabulador':            tabulador,
        'region':               region,
        'regiones':             regiones,
        'multiples_regiones':   len(regiones) > 1,
        'escuelas_asignadas':   escuelas_asignadas,
        'escuelas_disponibles': escuelas_disponibles,
        'total_asignadas':      len(escuelas_asignadas),
        'total_disponibles':    len(escuelas_disponibles),
    }
    return render(request, 'fluidez_octubre_2026/asignacion_tabulador.html', context)


#-------------------------------------------------CARGA DE APLICADORES-------------------------------------------------------------


@login_required
def buscar_persona_cuil(request):
    """
    Datos de una persona ya cargada como aplicador o tabulador (de cualquier
    región), para autocompletar los formularios de carga. Solo lectura.
    """
    if request.user.nivelacceso_id != "Regional":
        return JsonResponse({'ok': False, 'error': 'Sin permiso.'}, status=403)

    cuil = re.sub(r'\D', '', request.GET.get('cuil', ''))
    if len(cuil) != 11:
        return JsonResponse({'ok': False, 'error': 'CUIL inválido.'}, status=400)

    campos = ('apellido', 'nombre', 'correo', 'celular')
    aplicador = AplicadoresFluidezOctubre2026.objects.filter(cuil=cuil).values(*campos).first()
    tabulador = TabuladoresFluidezOctubre2026.objects.filter(cuil=cuil).values(*campos).first()
    persona = aplicador or tabulador
    if not persona:
        return JsonResponse({'ok': True, 'encontrado': False})

    # es_aplicador / es_tabulador: el formulario avisa si la persona ya está
    # cargada en su propia tabla (el alta la va a rechazar).
    return JsonResponse({
        'ok': True,
        'encontrado': True,
        'origen': 'aplicador' if aplicador else 'tabulador',
        'es_aplicador': bool(aplicador),
        'es_tabulador': bool(tabulador),
        **persona,
    })


# Región activa de la pantalla de aplicadores. Clave de sesión propia para no
# cambiar la región que tiene seleccionada la pantalla de tabuladores.
SESION_REGION_APLICADORES = 'apl_fluidez_oct26_region'
ERROR_SGE = 'No se pudieron consultar las secciones. Intentá de nuevo en unos minutos.'


def _region_aplicadores(request):
    """
    Regiones del regional y región activa. Devuelve (regiones, region) o
    (None, None) si el usuario no es Regional o no tiene región; en ese caso
    ya deja el mensaje de error cargado.
    """
    if request.user.nivelacceso_id != "Regional":
        messages.error(request, 'No tienes permiso para acceder a esta página.')
        return None, None

    regiones = obtener_regional(str(request.user.username))
    if not regiones:
        messages.error(request, 'No se encontró una región asignada a tu usuario.')
        return None, None
    if 'R.E. 2' in regiones:
        regiones = ['R.E. 2', 'SUB. R.E. 2-B']

    solicitada = request.GET.get('region')
    if solicitada and solicitada in regiones:
        request.session[SESION_REGION_APLICADORES] = solicitada

    region = request.session.get(SESION_REGION_APLICADORES)
    if not region or region not in regiones:
        region = regiones[0]
        request.session[SESION_REGION_APLICADORES] = region
    return regiones, region


def _redirect_gestion_aplicadores():
    return redirect(reverse('evaluaciones_educativas:fluidez_octubre_2026:gestion_aplicadores'))


@login_required
def gestion_aplicadores(request):
    """Lista de aplicadores de la región y avance de secciones asignadas."""
    regiones, region = _region_aplicadores(request)
    if not region:
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    aplicadores = (
        AplicadoresFluidezOctubre2026.objects
        .filter(region=region)
        .annotate(cant_secciones=Count('secciones'))
        .order_by('apellido', 'nombre')
    )

    # Total de secciones de la región (SGE) y cuántas ya tienen aplicador.
    total_secciones = None
    secciones_asignadas = None
    try:
        ids_region = [s['id_seccion'] for s in secciones_region(region)]
        total_secciones = len(ids_region)
        secciones_asignadas = SeccionesAplicadorFluidezOctubre2026.objects.filter(
            id_seccion__in=ids_region
        ).count()
    except DatabaseError:
        logger.exception('Error consultando secciones en sge_nacion (región %s)', region)
        messages.error(request, ERROR_SGE)

    context = {
        'region':              region,
        'regiones':            regiones,
        'multiples_regiones':  len(regiones) > 1,
        'aplicadores':         aplicadores,
        'total_aplicadores':   aplicadores.count(),
        'total_secciones':     total_secciones,
        'secciones_asignadas': secciones_asignadas,
    }
    return render(request, 'fluidez_octubre_2026/gestion_aplicadores.html', context)


@login_required
def carga_aplicador(request, cuil_aplicador=None):
    """Alta de aplicador, o edición si viene `cuil_aplicador`."""
    regiones, region = _region_aplicadores(request)
    if not region:
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    instancia = None
    if cuil_aplicador:
        instancia = get_object_or_404(AplicadoresFluidezOctubre2026, cuil=cuil_aplicador, region=region)

    form = AplicadorFluidezOctubreForm(request.POST or None, instance=instancia)
    if request.method == 'POST' and form.is_valid():
        aplicador = form.save(commit=False)
        if not instancia:
            aplicador.region = region
        aplicador.save()
        accion = 'actualizado' if instancia else 'cargado'
        messages.success(request, f'Aplicador {aplicador.apellido}, {aplicador.nombre} (CUIL: {aplicador.cuil}) {accion} correctamente.')
        if instancia:
            return _redirect_gestion_aplicadores()
        # Recién creado: pasar directo a asignarle secciones.
        return redirect(reverse(
            'evaluaciones_educativas:fluidez_octubre_2026:asignacion_aplicador',
            kwargs={'cuil_aplicador': aplicador.cuil},
        ))

    context = {
        'form':       form,
        'region':     region,
        'es_edicion': bool(instancia),
    }
    return render(request, 'fluidez_octubre_2026/carga_aplicador.html', context)


@login_required
@require_POST
def eliminar_aplicador(request, cuil_aplicador):
    """Elimina un aplicador de la región activa; sus secciones quedan libres."""
    regiones, region = _region_aplicadores(request)
    if not region:
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    aplicador = AplicadoresFluidezOctubre2026.objects.filter(cuil=cuil_aplicador, region=region).first()
    if aplicador:
        nombre_completo = f'{aplicador.apellido}, {aplicador.nombre}'
        aplicador.delete()  # CASCADE: borra también sus secciones
        messages.success(request, f'Aplicador {nombre_completo} (CUIL: {cuil_aplicador}) eliminado. Sus secciones quedaron disponibles.')
    else:
        messages.warning(request, 'El aplicador ya no existe.')
    return _redirect_gestion_aplicadores()


@login_required
def asignacion_aplicador(request, cuil_aplicador):
    """Asignar y quitar secciones de 2º y 3º grado a un aplicador."""
    regiones, region = _region_aplicadores(request)
    if not region:
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    aplicador = get_object_or_404(AplicadoresFluidezOctubre2026, cuil=cuil_aplicador, region=region)
    es_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'

    # ── POST: agregar o quitar una sección ───────────────────────────
    if request.method == 'POST':
        accion = request.POST.get('accion')
        try:
            id_seccion = int(request.POST.get('id_seccion', ''))
        except ValueError:
            id_seccion = None

        nivel, texto = 'error', 'Acción no válida.'

        if accion == 'agregar' and id_seccion:
            # Se vuelve a leer la sección de SGE: el front solo manda el id,
            # así no se puede asignar una sección de otra región o grado.
            try:
                encontradas = secciones_region(region, id_seccion=id_seccion)
            except DatabaseError:
                logger.exception('Error consultando la sección %s en sge_nacion', id_seccion)
                encontradas = None

            if encontradas is None:
                nivel, texto = 'error', ERROR_SGE
            elif not encontradas:
                nivel, texto = 'error', 'La sección no es de 2º o 3º grado de una escuela de tu región.'
            else:
                s = encontradas[0]
                try:
                    with transaction.atomic(using=SeccionesAplicadorFluidezOctubre2026.objects.db):
                        SeccionesAplicadorFluidezOctubre2026.objects.create(
                            aplicador=aplicador,
                            id_seccion=s['id_seccion'],
                            id_institucion=s['id_institucion'],
                            cueanexo=s['cueanexo'],
                            c_grado_nivel_servicio=s['c_grado_nivel_servicio'],
                            anio_grado=s['anio_grado'],
                            turno=s['turno'],
                            ciclo_lectivo=str(s['ciclo_lectivo']) if s['ciclo_lectivo'] is not None else None,
                        )
                    nivel, texto = 'success', f"{s['escuela']}: {s['label']} asignada correctamente."
                    if s['plurigrado']:
                        texto += (
                            f" Es una sección plurigrado: {aplicador.apellido}, {aplicador.nombre} "
                            f"queda a cargo de los alumnos de {s['grados_corto']} grado de la Sección {s['nombre_seccion']}."
                        )
                except IntegrityError:
                    # id_seccion es único: la sección ya tiene aplicador (incluye
                    # dos regionales asignándola al mismo tiempo).
                    nivel, texto = 'error', 'Esa sección ya tiene un aplicador asignado.'

        elif accion == 'quitar' and id_seccion:
            asignada = SeccionesAplicadorFluidezOctubre2026.objects.filter(
                aplicador=aplicador, id_seccion=id_seccion
            ).first()
            if asignada:
                asignada.delete()
                texto = 'Sección quitada correctamente.'
                if '/' in (asignada.anio_grado or ''):
                    texto += f' Era plurigrado: quedaron sin aplicador los alumnos de {asignada.anio_grado} de esa sección.'
                nivel = 'success'
            else:
                nivel, texto = 'warning', 'La sección no estaba asignada a este aplicador.'

        if es_ajax:
            return JsonResponse({
                'ok': nivel in ('success', 'warning'),
                'nivel': nivel,
                'mensaje': texto,
                'id_seccion': id_seccion,
                'accion': accion,
            }, status=200 if nivel != 'error' else 409)

        getattr(messages, nivel)(request, texto)
        return redirect(reverse(
            'evaluaciones_educativas:fluidez_octubre_2026:asignacion_aplicador',
            kwargs={'cuil_aplicador': cuil_aplicador},
        ))

    # ── GET: secciones asignadas y disponibles ───────────────────────
    asignadas_qs = list(aplicador.secciones.all())

    error_sge = False
    try:
        todas = secciones_region(region)
    except DatabaseError:
        logger.exception('Error consultando secciones en sge_nacion (región %s)', region)
        todas = []
        error_sge = True
        messages.error(request, ERROR_SGE)

    por_id = {s['id_seccion']: s for s in todas}
    ocupadas = set(
        SeccionesAplicadorFluidezOctubre2026.objects
        .filter(id_seccion__in=list(por_id))
        .values_list('id_seccion', flat=True)
    )

    # Asignadas: datos de SGE para mostrar escuela y nombre de sección; si la
    # sección ya no figura en la vista, se muestra lo guardado.
    secciones_asignadas = []
    for sec in asignadas_qs:
        datos = por_id.get(sec.id_seccion)
        secciones_asignadas.append(datos or {
            'id_seccion': sec.id_seccion,
            'cueanexo': sec.cueanexo,
            'escuela': f'(Escuela {sec.cueanexo})',
            'label': f"{sec.anio_grado or ''} — Sección {sec.id_seccion} — {sec.turno or 'Sin turno'}",
            'plurigrado': '/' in (sec.anio_grado or ''),
            'grados_corto': sec.anio_grado or '',
            'nombre_seccion': sec.id_seccion,
        })
    secciones_asignadas.sort(key=lambda s: (s.get('escuela') or '', s.get('label') or ''))

    secciones_disponibles = [s for s in todas if s['id_seccion'] not in ocupadas]

    context = {
        'aplicador':             aplicador,
        'region':                region,
        'secciones_asignadas':   secciones_asignadas,
        'secciones_disponibles': secciones_disponibles,
        'total_asignadas':       len(secciones_asignadas),
        'total_disponibles':     len(secciones_disponibles),
        'error_sge':             error_sge,
        'hay_plurigrado':        any(s.get('plurigrado') for s in secciones_asignadas + secciones_disponibles),
    }
    return render(request, 'fluidez_octubre_2026/asignacion_aplicador.html', context)


#---------------------------------descargar excel--------------------------
@login_required
def descargar_excel_personal_fluidez_octubre(request):
    """Excel con las secciones de la región y su tabulador y aplicador asignados."""
    nivel_acceso = request.user.nivelacceso_id
    if nivel_acceso != "Regional":
        messages.error(request, 'No tienes permiso para acceder a esta página.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    # Región: la que tiene seleccionada la pantalla (?region=), siempre que
    # sea una de las regiones del usuario. Mismo criterio que las gestiones.
    regiones = obtener_regional(str(request.user.username))
    if not regiones:
        messages.error(request, 'No se encontró una región asignada a tu usuario.')
        return redirect(reverse('evaluaciones_educativas:dashboard'))
    if 'R.E. 2' in regiones:
        regiones = ['R.E. 2', 'SUB. R.E. 2-B']
    region = request.GET.get('region')
    if region not in regiones:
        region = regiones[0]

    # Secciones de SGE de la región (una sola consulta), indexadas por cueanexo.
    try:
        secciones = secciones_region(region)
    except DatabaseError:
        logger.exception('Error consultando secciones en sge_nacion (región %s)', region)
        messages.error(request, ERROR_SGE)
        return redirect(reverse('evaluaciones_educativas:dashboard'))

    secciones_por_cueanexo = {}
    for s in secciones:
        secciones_por_cueanexo.setdefault(s['cueanexo'], []).append(s)

    # Aplicador de cada sección asignada, indexado por id_seccion.
    aplicador_por_seccion = {
        sa.id_seccion: sa.aplicador
        for sa in SeccionesAplicadorFluidezOctubre2026.objects
        .filter(aplicador__region=region)
        .select_related('aplicador')
    }

    def nombre_completo(persona):
        if not persona:
            return ''
        return f"{persona.apellido or ''}, {persona.nombre or ''}"

    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Datos Personal Fluidez 2026')
    ws.append([
        'Regional', 'Cueanexo', 'Escuela', 'Grados', 'Sección', 'Turno',
        'Apellido y Nombres Tabulador', 'CUIL Tabulador',
        'Apellido y Nombres Aplicador', 'CUIL Aplicador',
    ])

    tabuladores = TabuladoresFluidezOctubre2026.objects.filter(region=region).order_by('apellido', 'nombre')
    for tab in tabuladores:
        for cueanexo in tab.lista_cueanexos or []:
            for s in secciones_por_cueanexo.get(str(cueanexo).strip(), []):
                aplicador = aplicador_por_seccion.get(s['id_seccion'])
                ws.append([
                    region,
                    s['cueanexo'],
                    s['escuela'],
                    s['anio_grado'],
                    str(s['nombre_seccion'] or ''),
                    s['turno'] or '',
                    nombre_completo(tab),
                    tab.cuil,
                    nombre_completo(aplicador),
                    aplicador.cuil if aplicador else '',
                ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = (
        f'attachment; filename="datos_personal_fluidez_2026_{datetime.now():%Y%m%d_%H%M}.xlsx"'
    )
    wb.save(response)
    return response